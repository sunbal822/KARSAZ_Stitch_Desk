"""Reads and updates the orders sheet, and checks the workload against deadlines."""
import datetime, math, os, threading
from quote import HEADERS

STATUSES = ["Quoted", "Pending", "In progress", "Done", "Cancelled"]
COMMITTED = ("pending", "in progress")        # work you have agreed to do (a plain quote is not counted)
_lock = threading.Lock()


def _date(s):
    try:
        return datetime.date.fromisoformat(str(s)[:10])
    except (TypeError, ValueError):
        return None


def read_orders(path):
    """List of dicts, newest first. 'row' is the line in the sheet; 'due' is a date or None."""
    if not path or not os.path.exists(path):
        return []
    from openpyxl import load_workbook
    with _lock:
        wb = load_workbook(path, read_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
        wb.close()
    if not rows or [str(c) for c in rows[0][:len(HEADERS)]] != HEADERS:
        return []
    out = []
    for i, r in enumerate(rows[1:], start=2):
        if not any(r):
            continue
        d = dict(zip(HEADERS, r))
        d["row"] = i
        d["due"] = _date(d.get("Due date"))
        d["status"] = (d.get("Status") or "Quoted")
        out.append(d)
    return out[::-1]


def update_status(path, row, status):
    from openpyxl import load_workbook
    if status not in STATUSES:
        return "Pick a status from the list."
    with _lock:
        try:
            wb = load_workbook(path)
            ws = wb.active
            ws.cell(row=int(row), column=HEADERS.index("Status") + 1, value=status)
            wb.save(path)
        except PermissionError:
            return "Close orders.xlsx in Excel, then try again."
    return ""


def is_committed(o):
    return str(o["status"]).strip().lower() in COMMITTED


def _hours(o):
    try:
        return float(o.get("Hours") or 0)
    except (TypeError, ValueError):
        return 0.0


def committed_hours(orders):
    return sum(_hours(o) for o in orders if is_committed(o))


def overdue(orders, today=None):
    today = today or datetime.date.today()
    return [o for o in orders if is_committed(o) and o["due"] and o["due"] < today]


def _at_risk(jobs, hpd, today):
    """jobs: [(key, hours, due or None)]. Works through them earliest deadline first;
    returns the keys that cannot be finished in time at hpd hours a day."""
    late, cum = set(), 0.0
    for key, h, due in sorted(jobs, key=lambda j: (j[2] is None, j[2] or datetime.date.max)):
        cum += h
        if due is not None and math.ceil(cum / hpd) > (due - today).days:
            late.add(key)
    return late


def queue_earliest(orders, new_hours, cfg, today=None):
    """Earliest ready date for a new order if the committed orders are done first."""
    today = today or datetime.date.today()
    s = cfg["schedule"]
    return today + datetime.timedelta(days=math.ceil((committed_hours(orders) + new_hours) / s["hours_per_day"]) + s["buffer_days"])


def workload_warnings(orders, new_hours, new_due, cfg, today=None):
    """Plain sentences about a new order: can it meet its date, and does it push other orders late?"""
    today = today or datetime.date.today()
    hpd = float(cfg["schedule"]["hours_per_day"]) or 1.0
    jobs = [(o["row"], _hours(o), o["due"]) for o in orders if is_committed(o) and _hours(o) > 0]
    before = _at_risk(jobs, hpd, today)
    after = _at_risk(jobs + [("new", new_hours, new_due)], hpd, today)
    msgs = []
    if "new" in after:
        ahead = sum(h for _, h, due in jobs if due is not None and due <= new_due)   # orders due before it come first
        need = math.ceil((ahead + new_hours) / hpd)
        msgs.append(f"Finishing this by {new_due:%d %b} would take about {need} days of crocheting at {hpd:g} h a day "
                    f"(counting your agreed orders due before it), but only {max((new_due - today).days, 0)} days are left.")
    by_row = {o["row"]: o for o in orders}
    for r in sorted(after - before - {"new"}, key=lambda r: by_row[r]["due"]):
        o = by_row[r]
        msgs.append(f"This would push {o.get('Customer') or 'an order'}'s {o.get('Item')} (due {o['due']:%d %b}) past its date.")
    return msgs


def summary(orders, today=None):
    today = today or datetime.date.today()
    counts = {s: 0 for s in STATUSES}
    for o in orders:
        counts[o["status"] if o["status"] in counts else "Quoted"] += 1
    month = today.strftime("%Y-%m")
    booked = sum((o.get("Price (Rs)") or 0) for o in orders
                 if str(o.get("Date", ""))[:7] == month and str(o["status"]).lower() in ("pending", "in progress", "done"))
    return counts, booked


def late_orders(orders, cfg, today=None):
    """Agreed orders that cannot be finished by their date at your daily hours (earliest deadline first)."""
    today = today or datetime.date.today()
    hpd = float(cfg["schedule"]["hours_per_day"]) or 1.0
    jobs = [(o["row"], _hours(o), o["due"]) for o in orders if is_committed(o) and _hours(o) > 0]
    keys = _at_risk(jobs, hpd, today)
    return [o for o in orders if o["row"] in keys]