"""Crochet quote + order log.
Usage:
  python quote.py "romper_basic" --age 24 --colors 3 --changes 1
  python quote.py "lily" --log --customer "Ayesha"
Numbers live in pricing_config.json / catalog.json; no code changes needed.
"""
import argparse, datetime, json, math, os

HERE = os.path.dirname(os.path.abspath(__file__))


def _find(name):
    """File in this folder, matched ignoring upper/lower case (pricing_config.JSON still works on a Linux server)."""
    p = os.path.join(HERE, name)
    if os.path.exists(p): return p
    for f in os.listdir(HERE):
        if f.lower() == name.lower(): return os.path.join(HERE, f)
    return p


def load_cfg(path=None):
    with open(path or CFG_PATH, encoding="utf-8") as f:
        return json.load(f)


CFG_PATH = _find("pricing_config.json")
CFG = load_cfg(CFG_PATH)
CAT = json.load(open(_find("catalog.json"), encoding="utf-8"))
LOG = os.path.join(HERE, "orders.xlsx")

HEADERS = ["Date", "Customer", "Item", "Qty", "Colors requested", "Changes", "Size (months from>to)",
           "Grams", "Hours", "Yarn cost (Rs)", "Total cost (Rs)", "Price (Rs)", "Status", "Due date", "Notes"]


def _interp(points, x):
    """points = [[x, v1, v2...], ...] sorted by x. Returns the interpolated values for x."""
    if x <= points[0][0]: return points[0][1:]
    if x >= points[-1][0]: return points[-1][1:]
    for p0, p1 in zip(points, points[1:]):
        if p0[0] <= x <= p1[0]:
            t = (x - p0[0]) / (p1[0] - p0[0])
            return [a + (b - a) * t for a, b in zip(p0[1:], p1[1:])]


def size_factor(age, cfg=None):
    return _interp((cfg or CFG)["age_size_factors_months"], age)[0]


def size_chart(age, cfg=None):
    h, c = _interp((cfg or CFG)["size_chart_months"], age)
    return {"height_cm": round(h), "chest_cm": round(c)}


def find_item(key):
    k = key.lower()
    if k in CAT: return k
    for cid, v in CAT.items():
        if v.get("name", "").lower() == k: return cid
    raise KeyError(f"'{key}' is not in catalog.json")


def quote(item, colors=None, age=None, qty=1, n_changes=0, cfg=None):
    CFG = cfg or globals()["CFG"]   # a user's own prices, or the shop defaults
    item = find_item(item)
    it = CAT[item]
    colors = colors or it["colors"]
    base_age = it.get("base_age_months")
    wearable = base_age is not None
    factor = size_factor(age, CFG) / size_factor(base_age, CFG) if (age is not None and wearable) else 1.0
    cx = CFG["complexity"]
    complexity = min(1 + cx["extra_change_multiplier"] * n_changes, cx["max"])
    grams = it["grams"] * factor
    hours = it["hours"] * factor * complexity
    yarn = grams / 100 * CFG["yarn_cost_per_100g"]
    labor = hours * CFG["hourly_rate"]
    extra = CFG["extra_color_charge"] * max(colors - 1, 0)
    cost = yarn + labor + extra
    markup = max(CFG["category_markup"].get(it["category"], CFG["default_markup"]), CFG["min_markup"])
    r = CFG["round_to"]
    unit = round(cost * markup / r) * r
    chart_age = age if age is not None else base_age
    return dict(item=it.get("name", item), item_id=item, category=it["category"], qty=qty, colors=colors,
                size_factor=round(factor, 2), complexity=round(complexity, 2),
                grams=round(grams), hours=round(hours, 1), yarn_cost=round(yarn), cost=round(cost),
                markup=markup, unit_price=unit, total_price=unit * qty,
                size_from=base_age, size_to=age if wearable else None,
                size_cm=size_chart(chart_age, CFG) if (wearable and chart_age is not None) else None)


def due_date(total_hours, start=None, cfg=None):
    s = (cfg or CFG)["schedule"]
    start = start or datetime.date.today()
    return start + datetime.timedelta(days=math.ceil(total_hours / s["hours_per_day"]) + s["buffer_days"])


def log_order(q, customer="", colors_text="", changes_text="", due=None, notes="", status=None, path=None, cfg=None):
    from openpyxl import Workbook, load_workbook
    LOG = path or globals()["LOG"]
    CFG = cfg or globals()["CFG"]
    os.makedirs(os.path.dirname(LOG) or ".", exist_ok=True)
    if os.path.exists(LOG):
        wb = load_workbook(LOG)
        ws = wb.active
        if [c.value for c in ws[1]] != HEADERS:  # old layout: keep it, start a new log
            wb.close()
            os.replace(LOG, LOG.replace(".xlsx", f"_old_{datetime.date.today()}.xlsx"))
            wb = None
    else:
        wb = None
    if wb is None:
        wb = Workbook(); ws = wb.active; ws.title = "Orders"; ws.append(HEADERS)
    size = f"{q['size_from']:g}>{q['size_to']:g}" if q["size_from"] is not None and q["size_to"] is not None else ""
    n = q["qty"]
    ws.append([datetime.date.today().isoformat(), customer, q["item"], n, colors_text, changes_text, size,
               q["grams"] * n, round(q["hours"] * n, 1), q["yarn_cost"] * n, q["cost"] * n, q["total_price"],
               status or CFG["status_default"], str(due) if due else "", notes])
    try:
        wb.save(LOG)
    except PermissionError:
        raise RuntimeError("Close orders.xlsx in Excel, then try again.")


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("item"); p.add_argument("--colors", type=int); p.add_argument("--age", type=float)
    p.add_argument("--qty", type=int, default=1); p.add_argument("--changes", type=int, default=0)
    p.add_argument("--log", action="store_true")
    p.add_argument("--customer", default=""); p.add_argument("--notes", default="")
    a = p.parse_args()
    q = quote(a.item, a.colors, a.age, a.qty, a.changes)
    for k, v in q.items(): print(f"{k:12s}{v}")
    due = due_date(q["hours"] * q["qty"])
    print(f"{'due':12s}{due}")
    if a.log:
        log_order(q, a.customer, due=due, notes=a.notes); print("logged to", LOG)